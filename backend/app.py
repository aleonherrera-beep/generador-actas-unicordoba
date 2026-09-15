import io
import os
import re
import json
import tempfile
import subprocess
import logging
import traceback
import asyncio
from pathlib import Path
from typing import List, Optional, Dict, Any

import fitz
from docx import Document
from docx.shared import Pt
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse

BASE_DIR = Path(__file__).resolve().parent
ASSETS = BASE_DIR / "assets"
TEMPLATE_PATH = ASSETS / "template_acta.docx"
STYLE_PATH = BASE_DIR / "style_reference.txt"

logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
logger = logging.getLogger("actas")

app = FastAPI(title="Generador de Actas Unicórdoba", version="0.2.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "https://aleonherrera-beep.github.io",
        "http://localhost",
        "http://127.0.0.1",
    ],
    allow_origin_regex=r"https://.*\.github\.io",
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)

@app.middleware("http")
async def log_requests(request, call_next):
    logger.info("%s %s", request.method, request.url.path)
    try:
        response = await call_next(request)
        logger.info("%s %s -> %s", request.method, request.url.path, response.status_code)
        return response
    except Exception:
        logger.error("Error no controlado en %s %s\n%s", request.method, request.url.path, traceback.format_exc())
        raise

# ---------------------------
# Utilidades de documentos
# ---------------------------

def extract_pdf_text(data: bytes) -> str:
    doc = fitz.open(stream=data, filetype="pdf")
    parts = []
    for page in doc:
        parts.append(page.get_text("text"))
    text = "\n".join(parts).strip()
    return text


def extract_docx_text(data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        f.write(data)
        path = f.name
    try:
        d = Document(path)
        chunks = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                chunks.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(chunks)
    finally:
        os.unlink(path)


def normalize_spaces(s: str) -> str:
    return re.sub(r"[ \t]+", " ", s).strip()


def clean_person_name(name: str) -> str:
    name = normalize_spaces(name)
    return name.strip(" :-–—,.;")


def parse_citation(text: str) -> Dict[str, Any]:
    # Preserva líneas útiles y elimina ruido repetitivo.
    raw_lines = [normalize_spaces(x) for x in text.replace("\r", "\n").split("\n")]
    lines = [x for x in raw_lines if x]
    joined = "\n".join(lines)

    result: Dict[str, Any] = {
        "committee": "Comité de Postgrados",
        "subject": "",
        "date": "",
        "time": "",
        "place": "",
        "people": [],
        "agenda": [],
        "raw_text": text,
    }

    # Datos de reunión
    for key, pattern in {
        "subject": r"ASUNTO\s*:\s*(.+)",
        "date": r"FECHA\s*:\s*(.+)",
        "time": r"HORA\s*:\s*(.+)",
        "place": r"LUGAR\s*:\s*(.+)",
    }.items():
        m = re.search(pattern, joined, re.I)
        if m:
            result[key] = normalize_spaces(m.group(1))

    # Extrae bloques Para e Invitados. Trabajamos por líneas porque los PDF suelen partir cargos.
    def collect_block(start_label: str, end_labels: List[str]) -> List[str]:
        start = None
        for i, line in enumerate(lines):
            if re.match(rf"^{re.escape(start_label)}\s*:?\s*", line, re.I):
                start = i
                break
        if start is None:
            return []
        out = []
        first = re.sub(rf"^{re.escape(start_label)}\s*:?\s*", "", lines[start], flags=re.I).strip()
        if first:
            out.append(first)
        for line in lines[start+1:]:
            if any(re.match(rf"^{re.escape(lbl)}\s*:?", line, re.I) for lbl in end_labels):
                break
            out.append(line)
        return out

    members_lines = collect_block("Para", ["Invitados", "ASUNTO", "Estimados"])
    guests_lines = collect_block("Invitados", ["ASUNTO", "Estimados"])

    # Une líneas de cargo continuadas. Nueva persona suele comenzar con nombre en mayúsculas o patrón "Nombre, cargo".
    def parse_people(block: List[str], category: str):
        items = []
        current = ""
        for line in block:
            if not line:
                continue
            # Si línea contiene coma, suele empezar una persona.
            if "," in line and (current == "" or re.match(r"^[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ .'-]{3,},", line)):
                if current:
                    items.append(current)
                current = line
            else:
                # Detectar inicio por varias palabras en mayúscula seguido de cargo.
                m = re.match(r"^([A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ .'-]{4,})(,|\s{2,})(.*)$", line)
                if m and current:
                    items.append(current)
                    current = line
                elif current:
                    current += " " + line
                else:
                    current = line
        if current:
            items.append(current)

        people = []
        for item in items:
            item = normalize_spaces(item)
            if not item:
                continue
            if "," in item:
                name, role = item.split(",", 1)
            else:
                # Aproximación: primeras 2-4 palabras con capitalización alta.
                toks = item.split()
                cut = min(4, len(toks))
                name = " ".join(toks[:cut])
                role = " ".join(toks[cut:])
            name = clean_person_name(name.title() if name.isupper() else name)
            role = normalize_spaces(role)
            if len(name) >= 4:
                people.append({"name": name, "role": role, "category": category})
        return people

    people = parse_people(members_lines, "Miembro") + parse_people(guests_lines, "Invitado")
    # Deduplicación
    seen = set()
    dedup = []
    for p in people:
        key = re.sub(r"\W+", "", p["name"].lower())
        if key and key not in seen:
            seen.add(key)
            dedup.append(p)
    result["people"] = dedup

    # Orden del día: toma líneas posteriores al encabezado hasta pie institucional.
    agenda_idx = None
    for i, line in enumerate(lines):
        if re.search(r"ORDEN DEL D[IÍ]A", line, re.I):
            agenda_idx = i
            break
    if agenda_idx is not None:
        agenda = []
        for line in lines[agenda_idx+1:]:
            if re.search(r"Renovaci[oó]n de la acreditaci[oó]n|Certificados en:|PBX:|Se agradece", line, re.I):
                break
            line = re.sub(r"^[•\-–—\d\.\)\s]+", "", line).strip()
            if not line:
                continue
            if agenda and line[0].islower():
                agenda[-1] += " " + line
            elif agenda and len(line.split()) <= 5 and not line.endswith("."):
                agenda[-1] += " " + line
            else:
                agenda.append(line)
        result["agenda"] = [normalize_spaces(x) for x in agenda if len(x) > 3]

    return result


# ---------------------------
# Transcripción
# ---------------------------

_WHISPER = None
_WHISPER_NAME = None

def transcribe_audio(path: str, model_name: str = "small") -> str:
    global _WHISPER, _WHISPER_NAME
    from faster_whisper import WhisperModel
    import torch
    device = "cuda" if torch.cuda.is_available() else "cpu"
    compute_type = "float16" if device == "cuda" else "int8"
    if _WHISPER is None or _WHISPER_NAME != model_name:
        _WHISPER = WhisperModel(model_name, device=device, compute_type=compute_type)
        _WHISPER_NAME = model_name
    segments, info = _WHISPER.transcribe(
        path,
        language="es",
        vad_filter=True,
        beam_size=5,
    )
    out = []
    for seg in segments:
        start = int(seg.start)
        h, rem = divmod(start, 3600)
        m, s = divmod(rem, 60)
        out.append(f"[{h:02d}:{m:02d}:{s:02d}] {seg.text.strip()}")
    return "\n".join(out)


# ---------------------------
# Modelo de redacción local
# ---------------------------

_MODEL = None
_TOKENIZER = None
_MODEL_NAME = None

def load_llm(model_name: str = "Qwen/Qwen2.5-1.5B-Instruct"):
    """Carga un modelo más ligero y estable para Colab T4.

    Se evita bitsandbytes/4-bit porque en Colab suele generar conflictos de
    dependencias. En una T4, este modelo cabe cómodamente en FP16.
    """
    global _MODEL, _TOKENIZER, _MODEL_NAME
    if _MODEL is not None and _MODEL_NAME == model_name:
        return _TOKENIZER, _MODEL

    import torch
    from transformers import AutoTokenizer, AutoModelForCausalLM

    logger.info("Cargando modelo de redacción: %s", model_name)
    tok = AutoTokenizer.from_pretrained(model_name, use_fast=True)

    if torch.cuda.is_available():
        model = AutoModelForCausalLM.from_pretrained(
            model_name,
            torch_dtype=torch.float16,
            device_map="auto",
            low_cpu_mem_usage=True,
        )
    else:
        model = AutoModelForCausalLM.from_pretrained(
            model_name,
            torch_dtype=torch.float32,
            low_cpu_mem_usage=True,
        )

    model.eval()
    _TOKENIZER, _MODEL, _MODEL_NAME = tok, model, model_name
    logger.info("Modelo de redacción listo")
    return tok, model


def llm_generate(messages: List[Dict[str, str]], max_new_tokens: int = 1400) -> str:
    import torch

    tok, model = load_llm()
    prompt = tok.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
    inputs = tok(prompt, return_tensors="pt", truncation=True, max_length=12000)

    try:
        dev = next(model.parameters()).device
        inputs = {k: v.to(dev) for k, v in inputs.items()}
    except Exception:
        pass

    logger.info("Generando texto con %s tokens de entrada", inputs["input_ids"].shape[1])
    with torch.inference_mode():
        outputs = model.generate(
            **inputs,
            max_new_tokens=max_new_tokens,
            do_sample=False,
            repetition_penalty=1.04,
            pad_token_id=tok.eos_token_id,
        )

    generated = outputs[0][inputs["input_ids"].shape[1]:]
    return tok.decode(generated, skip_special_tokens=True).strip()

def split_text(text: str, max_chars: int = 9000) -> List[str]:
    if len(text) <= max_chars:
        return [text]
    chunks = []
    current = []
    size = 0
    for p in re.split(r"\n+", text):
        if size + len(p) > max_chars and current:
            chunks.append("\n".join(current))
            current, size = [], 0
        current.append(p)
        size += len(p) + 1
    if current:
        chunks.append("\n".join(current))
    return chunks


def summarize_transcript(transcript: str, meeting: Dict[str, Any]) -> str:
    style = STYLE_PATH.read_text(encoding="utf-8")
    chunks = split_text(transcript, 6000)
    notes = []
    agenda = "\n".join(f"- {x}" for x in meeting.get("agenda", []))
    for idx, chunk in enumerate(chunks, 1):
        prompt = f"""
Eres secretario técnico de un comité universitario. Extrae NOTAS FÁCTICAS de este fragmento de transcripción.
No redactes todavía el acta final. No inventes. Conserva nombres, cifras, fechas, decisiones, responsables, solicitudes, recomendaciones y asuntos discutidos.
Relaciona, cuando sea posible, cada nota con el orden del día.

ORDEN DEL DÍA:
{agenda}

FRAGMENTO {idx}/{len(chunks)}:
{chunk}
"""
        notes.append(llm_generate([
            {"role": "system", "content": "Extraes información fiel de reuniones académicas. Nunca inventas hechos."},
            {"role": "user", "content": prompt},
        ], max_new_tokens=1200))
    return "\n\n".join(notes)


def extract_json(text: str) -> Dict[str, Any]:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    start = text.find("{")
    end = text.rfind("}")
    if start >= 0 and end > start:
        text = text[start:end+1]
    return json.loads(text)


def draft_minutes(transcript: str, meeting: Dict[str, Any], attendance: List[Dict[str, Any]]) -> Dict[str, Any]:
    style = STYLE_PATH.read_text(encoding="utf-8")
    notes = summarize_transcript(transcript, meeting)
    agenda = "\n".join(f"{i+1}. {x}" for i, x in enumerate(meeting.get("agenda", [])))
    att = "\n".join(f"- {p.get('name')}: {p.get('status')} ({p.get('role','')})" for p in attendance)

    schema = {
        "development": "texto narrativo completo del desarrollo de la sesión, con subtítulos según temas",
        "approval_previous": "texto breve; si no hay evidencia, N.A.",
        "previous_commitments": [],
        "correspondence": [{"sender": "", "subject": "", "decision": ""}],
        "varios": "texto o N.A.",
        "commitments": [{"task": "", "responsible": "", "due": "", "verification": ""}],
        "decisions": [{"decision": "", "responsible": "", "due": ""}],
        "end_time": "",
        "next_session": {"date": "", "time": "", "place": ""}
    }

    prompt = f"""
Redacta el contenido de un acta institucional universitaria a partir EXCLUSIVAMENTE de las notas fieles de una reunión.
Sigue el estilo descrito y devuelve SOLO JSON válido, sin markdown.

ESTILO:
{style}

DATOS DE LA REUNIÓN:
{json.dumps(meeting, ensure_ascii=False)}

ASISTENCIA CONFIRMADA MANUALMENTE:
{att}

ORDEN DEL DÍA:
{agenda}

NOTAS EXTRAÍDAS DE LA TRANSCRIPCIÓN:
{notes}

REGLAS:
- No inventes decisiones, responsables, fechas, cifras ni intervenciones.
- Redacta en tercera persona y tono institucional.
- El desarrollo debe ser sustancial, no un resumen corto; debe conservar el contenido importante de cada tema.
- Organiza el desarrollo siguiendo el orden del día y usando subtítulos claros.
- Si lectura de correspondencia contiene solicitudes y decisiones, sepáralas en el arreglo correspondence.
- Si no hay evidencia de un campo, usa "N.A." o arreglo vacío.
- No incluyas firmas inventadas.
- JSON objetivo con esta forma exacta:
{json.dumps(schema, ensure_ascii=False, indent=2)}
"""
    raw = llm_generate([
        {"role": "system", "content": "Redactas actas universitarias fieles, extensas, objetivas y estructuradas. Respondes en JSON válido."},
        {"role": "user", "content": prompt},
    ], max_new_tokens=2000)
    try:
        return extract_json(raw)
    except Exception:
        # Fallback: conserva el texto como desarrollo en vez de fallar por formato JSON.
        return {
            "development": raw,
            "approval_previous": "N.A.",
            "previous_commitments": [],
            "correspondence": [],
            "varios": "N.A.",
            "commitments": [],
            "decisions": [],
            "end_time": "",
            "next_session": {"date": "", "time": "", "place": ""},
        }


# ---------------------------
# Word institucional
# ---------------------------

from copy import deepcopy
from docx.oxml import OxmlElement


def set_cell_text(cell, text: str, bold: bool = False, size: Optional[int] = None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text or ""))
    run.bold = bold
    if size:
        run.font.size = Pt(size)


def remove_rows_after(table, keep_rows: int):
    while len(table.rows) > keep_rows:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)


def clone_row(table, source_index: int = -1):
    source = table.rows[source_index]._tr
    new_tr = deepcopy(source)
    table._tbl.append(new_tr)
    return table.rows[-1]


def fill_repeating_table(table, header_rows: int, items: List[Dict[str, Any]], columns: List[str], blank_when_empty=True):
    # Mantiene una fila modelo si existe.
    template_idx = header_rows if len(table.rows) > header_rows else len(table.rows)-1
    template = deepcopy(table.rows[template_idx]._tr) if template_idx >= 0 else None
    remove_rows_after(table, header_rows)
    if not items and blank_when_empty:
        items = [{c: "N.A." if i == 0 else "" for i, c in enumerate(columns)}]
    for item in items:
        if template is not None:
            table._tbl.append(deepcopy(template))
            row = table.rows[-1]
        else:
            row = table.add_row()
        for i, key in enumerate(columns):
            if i < len(row.cells):
                set_cell_text(row.cells[i], item.get(key, ""))


def build_docx(meeting: Dict[str, Any], attendance: List[Dict[str, Any]], draft: Dict[str, Any]) -> bytes:
    if not TEMPLATE_PATH.exists():
        raise RuntimeError("No se encontró la plantilla institucional.")
    doc = Document(TEMPLATE_PATH)
    tables = doc.tables
    if len(tables) < 14:
        raise RuntimeError("La plantilla no tiene la estructura esperada.")

    # 0 Comité
    set_cell_text(tables[0].cell(1, 0), meeting.get("committee") or "Comité de Postgrados")

    # 1 datos
    values = [
        meeting.get("act_number", ""), meeting.get("place", ""), meeting.get("date", ""),
        meeting.get("time", ""), draft.get("end_time", "")
    ]
    for i, v in enumerate(values):
        set_cell_text(tables[1].cell(1, i), v)

    # Miembros e invitados
    members = [p for p in attendance if p.get("category") == "Miembro"]
    guests = [p for p in attendance if p.get("category") != "Miembro" and p.get("status") == "Asistió"]

    t = tables[2]
    # conserva 3 filas encabezado, reconstruye personas
    template = deepcopy(t.rows[3]._tr) if len(t.rows) > 3 else None
    remove_rows_after(t, 3)
    for p in members:
        if template is not None:
            t._tbl.append(deepcopy(template))
            row = t.rows[-1]
        else:
            row = t.add_row()
        set_cell_text(row.cells[0], p.get("role", ""))
        set_cell_text(row.cells[1], p.get("name", ""))
        status = p.get("status", "")
        set_cell_text(row.cells[2], "X" if status == "Asistió" else "")
        set_cell_text(row.cells[3], "X" if status == "No asistió" else "")
        set_cell_text(row.cells[4], "X" if status == "Excusa" else "")

    tg = tables[3]
    template = deepcopy(tg.rows[2]._tr) if len(tg.rows) > 2 else None
    remove_rows_after(tg, 2)
    if not guests:
        guests = [{"name": "N.A.", "role": ""}]
    for p in guests:
        if template is not None:
            tg._tbl.append(deepcopy(template))
            row = tg.rows[-1]
        else:
            row = tg.add_row()
        set_cell_text(row.cells[0], p.get("name", ""))
        set_cell_text(row.cells[1], p.get("role", ""))

    # Orden del día
    agenda = meeting.get("agenda", [])
    agenda_text = "\n".join(f"{i+1}. {x}" for i, x in enumerate(["Verificación de quórum", "Desarrollo de la sesión"] + agenda))
    set_cell_text(tables[4].cell(1, 0), agenda_text)
    set_cell_text(tables[4].cell(2, 0), "MODIFICACIÓN AL ORDEN DEL DÍA: SI__ NO_X_")
    set_cell_text(tables[4].cell(3, 0), "NUEVO ORDEN DEL DÍA APROBADO\nN.A.")

    # Acta anterior / compromisos previos
    set_cell_text(tables[5].cell(1, 0), draft.get("approval_previous", "N.A."))
    fill_repeating_table(tables[6], 2, draft.get("previous_commitments", []), ["task", "responsible", "due", "verification"])

    # Desarrollo
    set_cell_text(tables[7].cell(1, 0), draft.get("development", ""))

    # Correspondencia
    fill_repeating_table(tables[8], 2, draft.get("correspondence", []), ["sender", "subject", "decision"])

    # Varios
    set_cell_text(tables[9].cell(1, 0), draft.get("varios", "N.A."))

    # Compromisos y decisiones
    fill_repeating_table(tables[10], 2, draft.get("commitments", []), ["task", "responsible", "due", "verification"])
    fill_repeating_table(tables[11], 2, draft.get("decisions", []), ["decision", "responsible", "due"])

    # Próxima sesión
    ns = draft.get("next_session") or {}
    set_cell_text(tables[12].cell(1, 1), ns.get("date", ""))
    set_cell_text(tables[12].cell(1, 2), ns.get("time", ""))
    set_cell_text(tables[12].cell(1, 3), ns.get("place", ""))

    # Firmas: se dejan nombres/cargos originales como plantilla solo si se especifican.
    # Para evitar inventar, pueden venir en meeting.
    if meeting.get("president_name"):
        set_cell_text(tables[13].cell(1, 0), f"NOMBRE: {meeting['president_name']}")
    else:
        set_cell_text(tables[13].cell(1, 0), "NOMBRE:")
    if meeting.get("president_role"):
        set_cell_text(tables[13].cell(2, 0), f"CARGO: {meeting['president_role']}")
    else:
        set_cell_text(tables[13].cell(2, 0), "CARGO:")
    if meeting.get("secretary_name"):
        set_cell_text(tables[13].cell(1, 1), f"NOMBRE: {meeting['secretary_name']}")
    else:
        set_cell_text(tables[13].cell(1, 1), "NOMBRE:")
    if meeting.get("secretary_role"):
        set_cell_text(tables[13].cell(2, 1), f"CARGO: {meeting['secretary_role']}")
    else:
        set_cell_text(tables[13].cell(2, 1), "CARGO:")

    # Elimina el resumen adicional que existe al final del documento modelo.
    # En la plantilla suministrada, empieza por "Resumen Acta No.".
    remove = False
    for p in list(doc.paragraphs):
        if p.text.strip().startswith("Resumen Acta No."):
            remove = True
        if remove:
            el = p._element
            el.getparent().remove(el)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---------------------------
# API
# ---------------------------

@app.get("/health")
def health():
    import torch
    return {
        "ok": True,
        "service": "Generador de Actas Unicórdoba",
        "version": "0.2.0",
        "gpu": bool(torch.cuda.is_available()),
        "model_loaded": _MODEL is not None,
    }

@app.get("/warmup")
def warmup():
    """Precarga el modelo sin pasar por una petición de redacción larga."""
    try:
        load_llm()
        return {"ok": True, "model": _MODEL_NAME, "message": "Modelo listo"}
    except Exception as e:
        logger.error("Fallo al precargar modelo\n%s", traceback.format_exc())
        raise HTTPException(status_code=500, detail=f"No se pudo cargar el modelo: {e}")

@app.get("/diagnostics")
def diagnostics():
    import torch
    data = {
        "ok": True,
        "model_loaded": _MODEL is not None,
        "model_name": _MODEL_NAME,
        "cuda_available": torch.cuda.is_available(),
    }
    if torch.cuda.is_available():
        data["gpu_name"] = torch.cuda.get_device_name(0)
        data["gpu_memory_gb"] = round(torch.cuda.get_device_properties(0).total_memory / 1024**3, 2)
    return data

@app.post("/parse-citation")
async def api_parse_citation(file: UploadFile = File(...)):
    data = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".pdf"):
        text = extract_pdf_text(data)
    elif name.endswith(".docx"):
        text = extract_docx_text(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise HTTPException(400, "Formato de citación no soportado. Use PDF, DOCX o TXT.")
    parsed = parse_citation(text)
    return parsed

@app.post("/transcribe")
async def api_transcribe(file: UploadFile = File(...), model: str = Form("small")):
    suffix = Path(file.filename or "audio.mp3").suffix or ".mp3"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as f:
        f.write(await file.read())
        path = f.name
    try:
        text = transcribe_audio(path, model)
        return {"transcript": text}
    finally:
        try: os.unlink(path)
        except OSError: pass

@app.post("/read-transcript")
async def api_read_transcript(file: UploadFile = File(...)):
    data = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    elif name.endswith(".docx"):
        text = extract_docx_text(data)
    elif name.endswith(".pdf"):
        text = extract_pdf_text(data)
    else:
        raise HTTPException(400, "Use TXT, DOCX o PDF para la transcripción.")
    return {"transcript": text}

@app.post("/draft")
async def api_draft(payload: Dict[str, Any]):
    transcript = payload.get("transcript", "").strip()
    if not transcript:
        raise HTTPException(400, "La transcripción está vacía.")
    meeting = payload.get("meeting") or {}
    attendance = payload.get("attendance") or []

    logger.info("Inicio de redacción: %s caracteres de transcripción", len(transcript))
    try:
        result = await asyncio.to_thread(draft_minutes, transcript, meeting, attendance)
        logger.info("Redacción terminada correctamente")
        return result
    except Exception as e:
        logger.error("Fallo en /draft\n%s", traceback.format_exc())
        raise HTTPException(status_code=500, detail=f"Error al redactar el acta: {type(e).__name__}: {e}")

@app.post("/docx")
async def api_docx(payload: Dict[str, Any]):
    meeting = payload.get("meeting") or {}
    attendance = payload.get("attendance") or []
    draft = payload.get("draft") or {}
    data = build_docx(meeting, attendance, draft)
    filename = f"Acta_{meeting.get('act_number','Nueva')}.docx".replace(" ", "_")
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
